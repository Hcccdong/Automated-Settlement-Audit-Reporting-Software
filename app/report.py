# -*- coding: utf-8 -*-
"""报告生成引擎：在模板副本上做占位符替换，格式零改动。
- 跨 run 替换，新文字继承匹配起始 run 的格式（下划线等不丢失）
- 多行字段按行拆段，新段落深拷贝占位段格式
- 支持正文、表格、页眉页脚中的占位符
"""
import copy
import os
import re
from datetime import datetime

from docx import Document
from docx.text.paragraph import Paragraph

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates", "结算审核模板_占位版.docx")
OUT_DIR = os.path.join(BASE, "outputs", "生成报告")

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]+\}\}")


_MC_URI = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _run_has_drawing(r):
    from docx.oxml.ns import qn
    el = r._element
    for tag in ("w:pict", "w:drawing", "w:object"):
        if el.find(".//" + qn(tag)) is not None:
            return True
    return el.find(".//{%s}AlternateContent" % _MC_URI) is not None


def _replace_run_aware(p, old, new):
    """跨 run 替换，保留全部 run 格式，新文字继承匹配起始 run 的格式。
    含图形（横线形状等）的 run 一律不动。"""
    runs = p.runs
    count = 0
    if old == new:
        return 0
    while True:
        full = "".join(r.text for r in runs)
        idx = full.find(old)
        if idx == -1:
            break
        end = idx + len(old)
        pos = 0
        for r in runs:
            rlen = len(r.text)
            rs, re_ = pos, pos + rlen
            pos = re_
            if re_ <= idx or rs >= end:
                continue
            if _run_has_drawing(r):
                continue
            local_s = max(idx - rs, 0)
            local_e = min(end - rs, rlen)
            if rs <= idx:
                r.text = r.text[:local_s] + new + (r.text[local_e:] if re_ >= end else "")
            else:
                r.text = r.text[local_e:] if re_ >= end else ""
        count += 1
    return count


def _set_text(p, text):
    """整段写入文本，保留第一个非图形 run 的格式；图形 run 不动。"""
    target = None
    for r in p.runs:
        if not _run_has_drawing(r):
            target = r
            break
    if target is None:
        p.add_run(text)
        return
    target.text = text
    for r in p.runs:
        if r._element is not target._element and not _run_has_drawing(r):
            r.text = ""


def _fill_paragraph(p, values):
    """替换段落内所有占位符；多行值拆段，新段落复制本段格式。"""
    text = "".join(r.text for r in p.runs)
    if "{{" not in text:
        return
    multiline = None
    for m in PLACEHOLDER_RE.finditer(text):
        key = m.group(0)[2:-2]
        if key not in values:
            continue
        val = str(values.get(key, m.group(0)))
        if "\n" in val:
            lines = val.split("\n")
            _replace_run_aware(p, m.group(0), lines[0])
            multiline = lines[1:]
        else:
            _replace_run_aware(p, m.group(0), val)
    if multiline:
        prev = p._element
        for line in multiline:
            new_el = copy.deepcopy(p._element)
            prev.addnext(new_el)
            prev = new_el
            new_p = Paragraph(new_el, p._parent)
            # 拆出的新段不携带图形 run（避免横线形状被复制到每一行）
            for r in list(new_p.runs):
                if _run_has_drawing(r):
                    r._element.getparent().remove(r._element)
            _set_text(new_p, line)


def _fill_container(container, values):
    for p in container.paragraphs:
        _fill_paragraph(p, values)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                _fill_container(cell, values)


def _to_number(s):
    s = str(s).replace(",", "").replace("，", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def validate(values):
    """金额勾稽校验，返回警告列表（不阻止生成）。"""
    warnings = []
    song = _to_number(values.get("送审金额", ""))
    shen = _to_number(values.get("审定金额", ""))
    zeng = _to_number(values.get("核增金额", ""))
    jian = _to_number(values.get("核减金额", ""))
    jing = _to_number(values.get("净核减金额", ""))
    if None in (song, shen, zeng, jian, jing):
        warnings.append("金额字段存在非数字内容，无法完成勾稽校验，请人工核对。")
        return warnings
    if abs((song - shen) - jing) > 0.005:
        warnings.append(f"勾稽不符：送审-审定={song - shen:,.0f}，与净核减 {jing:,.0f} 不一致。")
    if abs((jian - zeng) - jing) > 0.005:
        warnings.append(f"勾稽不符：核减-核增={jian - zeng:,.0f}，与净核减 {jing:,.0f} 不一致。")
    return warnings


def _all_paragraphs(doc):
    paras = list(doc.paragraphs)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header,
                   sec.first_page_footer, sec.even_page_header, sec.even_page_footer):
            try:
                paras.extend(hf.paragraphs)
            except Exception:
                pass
    return paras


def generate(values, template_path=None, out_dir=None):
    """生成报告，返回 (输出文件路径, 警告列表)。"""
    tpl = template_path or TEMPLATE
    out = out_dir or OUT_DIR
    os.makedirs(out, exist_ok=True)

    doc = Document(tpl)
    _fill_container(doc, values)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header,
                   sec.first_page_footer, sec.even_page_header, sec.even_page_footer):
            try:
                _fill_container(hf, values)
            except Exception:
                pass

    # 遗留占位符检查
    leftovers = []
    for p in _all_paragraphs(doc):
        for m in PLACEHOLDER_RE.finditer(p.text):
            if m.group(0)[2:-2] not in values:
                leftovers.append(m.group(0))

    name = str(values.get("工程名称", "未命名工程")).strip()
    safe = re.sub(r'[\\/:*?"<>|]', "_", name)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(out, f"结算审核报告_{safe}_{stamp}.docx")
    doc.save(path)

    warnings = validate(values)
    if leftovers:
        warnings.append("以下占位符未填写，已原样保留：" + "、".join(sorted(set(leftovers))))
    return path, warnings
