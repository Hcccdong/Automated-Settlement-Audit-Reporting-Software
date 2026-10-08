import tempfile
import unittest
from pathlib import Path

from docx import Document

import report


class ReportGenerationTests(unittest.TestCase):
    def test_missing_fields_remain_in_export_with_warning(self):
        with tempfile.TemporaryDirectory() as output:
            path, warnings = report.generate({"工程名称": "导出验证工程"}, out_dir=output)
            text = "\n".join(p.text for p in Document(path).paragraphs)
            self.assertIn("导出验证工程", text)
            self.assertIn("{{", text)
            self.assertTrue(any("占位符未填写" in warning for warning in warnings))
            self.assertTrue(Path(path).is_file())

    def test_identical_placeholder_replacement_returns(self):
        paragraph = Document().add_paragraph("{{工程名称}}")
        self.assertEqual(report._replace_run_aware(paragraph, "{{工程名称}}", "{{工程名称}}"), 0)
        self.assertEqual(paragraph.text, "{{工程名称}}")


if __name__ == "__main__":
    unittest.main()
