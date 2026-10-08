# -*- coding: utf-8 -*-
"""Excel 结算对比表解析与核减文字组稿。

解析：自动探测表头行与栏目（名称/单位工程/送审与审核的工程量/单价/金额/偏差），
      按单位工程分组提取明细行。
符号约定：净核减 = 送审 - 审核，正值为核减（文字称"负偏差"），负值为核增（"正偏差"）。
原因推断：仅工程量变→工程量有误；仅单价变→单价有误；两者皆变→工程量及单价有误；
      无量价依据→待确认（不凭金额断言原因）。
组稿：|偏差| 小于阈值（默认1000元）的条目忽略；未列条目与总额之间的差额，
      以"取费基数变化"说明段配平。
"""
import re

import openpyxl

# ---------------- 栏目关键词 ----------------
COL_KEYWORDS = {
    "name": ["项目名称", "单位工程名称", "名称", "事由", "事项", "内容", "项目"],
    "unit": ["单位工程"],
    "qty_s": ["送审工程量", "送审数量"],
    "qty_a": ["审核工程量", "审核数量"],
    "price_s": ["送审综合单价", "送审单价"],
    "price_a": ["审核综合单价", "审核单价"],
    "amt_s": ["施工方送审金额", "送审金额", "送审合价", "送审造价", "合同金额"],
    "amt_a": ["审核金额合计", "审核金额", "审核合价", "审核造价", "审定金额"],
    "diff": ["净核减", "核减额", "偏差", "核减（元）", "核减(元)"],
    "note": ["备注"],
}
# 长关键词优先匹配，避免"审核金额"吃掉"审核金额合计"
for k in COL_KEYWORDS:
    COL_KEYWORDS[k].sort(key=len, reverse=True)

CN_NUM = re.compile(r"^[一二三四五六七八九十百零0-9]+[、\.]?$")

# ---------------- 广联达"工程审核简明表"专用版式 ----------------
# 布局：col0=序号；col1-7 送审(编号/名称/项目特征/单位/数量/单价/合价)；
#       col8-14 审查(同构)；col15=核增；col16=核减
# 核增=正偏差(核增)，核减=负偏差(核减)；名称结合项目特征首行防歧义。
JM = {"no": 0, "s_code": 1, "s_name": 2, "s_feat": 3, "s_unit": 4,
      "s_qty": 5, "s_price": 6, "s_amt": 7,
      "a_code": 8, "a_name": 9, "a_feat": 10, "a_unit": 11,
      "a_qty": 12, "a_price": 13, "a_amt": 14, "zeng": 15, "jian": 16}


def _is_jianming(rows):
    """前 8 行内同时出现'项目特征'与'核增''核减'，判定为简明表版式。"""
    head = "\n".join(" ".join(str(c) for c in r if c is not None) for r in rows[:8])
    return "项目特征" in head and "核增" in head and "核减" in head and "送审" in head


def _feat_short(feat, limit=40):
    """项目特征取首行、去换行，超长截断。"""
    if not feat:
        return ""
    first = re.split(r"[\r\n]+", str(feat))[0].strip().strip("：:")
    return first[:limit]


def _parse_jianming_sheet(ws):
    rows = [list(r) + [None] * (17 - len(r)) for r in ws.iter_rows(values_only=True)]
    if not _is_jianming(rows):
        return None

    groups = []
    group_map = {}
    warnings = []

    def get_group(unit):
        if unit not in group_map:
            g = {"unit": unit, "items": []}
            group_map[unit] = g
            groups.append(g)
        return group_map[unit]

    default_unit = ws.title
    for row in rows:
        c0 = str(row[JM["no"]]).strip() if row[JM["no"]] is not None else ""
        joined = " ".join(str(c) for c in row[:3] if c is not None)
        # 工程名称行 → 单位工程分组
        if "工程名称" in joined:
            m = re.search(r"工程名称\s*[:：]\s*(.+?)(?:\s*第\d+页|$)", joined)
            if m:
                default_unit = m.group(1).strip()
            continue
        # 跳过页眉/表头/空行
        if c0 == "序号" or str(row[JM["s_code"]]).strip() == "编号" or "简明表" in joined:
            continue
        # 分部行（无序号、有编号如 0104）跳过
        if not c0 and row[JM["s_code"]] is not None:
            continue
        # 明细行：序号列为数字
        if not re.match(r"^\d+$", c0):
            continue

        name = str(row[JM["s_name"]] or "").strip()
        if not name:
            continue
        feat = _feat_short(row[JM["s_feat"]])
        disp = f"{name}（{feat}）" if feat else name

        qty_s = _to_num(row[JM["s_qty"]])
        qty_a = _to_num(row[JM["a_qty"]])
        price_s = _to_num(row[JM["s_price"]])
        price_a = _to_num(row[JM["a_price"]])
        amt_s = _to_num(row[JM["s_amt"]])
        amt_a = _to_num(row[JM["a_amt"]])
        zeng = _to_num(row[JM["zeng"]])
        jian = _to_num(row[JM["jian"]])

        if zeng is not None or jian is not None:
            net = (jian or 0.0) - (zeng or 0.0)
        elif amt_s is not None and amt_a is not None:
            net = amt_s - amt_a
        else:
            continue  # 无偏差、无金额，未发生变化
        if abs(net) < 1e-9:
            continue

        get_group(default_unit)["items"].append({
            "name": disp,
            "amt_s": amt_s, "amt_a": amt_a,
            "qty_s": qty_s, "qty_a": qty_a,
            "price_s": price_s, "price_a": price_a,
            "net": round(net, 2),
            "reason": infer_reason(qty_s, qty_a, price_s, price_a),
            "note": "",
            "excluded": False,
        })

    groups = [g for g in groups if g["items"]]
    if not groups:
        # 是简明表但没有任何偏差行（如空白模板）
        return {"sheet": ws.title, "columns": {"format": "简明表"}, "groups": [],
                "warnings": ["该表为审核简明表版式，但未检测到核增/核减数据（可能尚未填写调整内容）。"]}
    return {"sheet": ws.title, "columns": {"format": "简明表"}, "groups": groups, "warnings": warnings}



def _norm(v):
    return re.sub(r"[\s（）()：:]", "", str(v or ""))


def _to_num(v):
    if v is None or v == "":
        return None
    s = str(v).replace(",", "").replace("，", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def _fmt(n):
    """金额格式化：整数不带小数，否则两位；不带千分位（与报告习惯一致）。"""
    if n is None:
        return ""
    if abs(n - round(n)) < 0.005:
        return str(int(round(n)))
    return f"{n:.2f}"


def detect_header_row(rows, max_scan=12):
    """在前若干行中找关键词命中最多的行作为表头。"""
    best, best_hits = None, 0
    for i, row in enumerate(rows[:max_scan]):
        hits = 0
        for cell in row:
            c = _norm(cell)
            if not c:
                continue
            for kws in COL_KEYWORDS.values():
                if any(c == _norm(kw) or _norm(kw) in c for kw in kws):
                    hits += 1
                    break
        if hits > best_hits:
            best, best_hits = i, hits
    return best if best_hits >= 2 else None


def map_columns(header_row):
    """返回 {字段: 列下标}。先精确匹配，再包含匹配；名称列取最左。"""
    mapping = {}
    cells = [(_norm(c), idx) for idx, c in enumerate(header_row)]
    for field, kws in COL_KEYWORDS.items():
        nkw = [_norm(k) for k in kws]
        for kw in nkw:
            hit = next((idx for c, idx in cells if c == kw), None)
            if hit is None:
                hit = next((idx for c, idx in cells if kw and kw in c), None)
            if hit is not None:
                mapping[field] = hit
                break
    # 单位工程列不能与名称列相同（名称列优先保留更靠左的）
    if mapping.get("unit") == mapping.get("name"):
        del mapping["unit"]
    return mapping


def _parse_sheet(ws):
    """解析单个工作表，不合格返回 None。"""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return None
    hi = detect_header_row(rows)
    if hi is None:
        return None
    mapping = map_columns(rows[hi])
    if "name" not in mapping or ("amt_s" not in mapping and "amt_a" not in mapping and "diff" not in mapping):
        return None

    warnings = []
    groups = []
    current = {"unit": "（未分组）", "items": []}

    def get(row, field):
        idx = mapping.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    for row in rows[hi + 1:]:
        name = get(row, "name")
        unit_cell = get(row, "unit")
        # 单位工程列有值的行视为分组行（名称列可空）
        if unit_cell is not None and str(unit_cell).strip():
            gname = str(unit_cell).strip()
            amt_s0 = _to_num(get(row, "amt_s"))
            if amt_s0 is not None or _to_num(get(row, "amt_a")) is not None or _to_num(get(row, "diff")) is not None:
                warnings.append(f"分组行“{gname}”带有汇总金额，已按分组标题处理、未计入明细，"
                                "配平时请以其下明细合计为准。")
            current = {"unit": gname, "items": []}
            groups.append(current)
            continue
        if name is None or str(name).strip() == "":
            continue
        name = str(name).strip()
        if re.match(r"^-?[\d\.,]+$", name):
            continue  # 纯数字名称多为表外演算数据，忽略
        amt_s = _to_num(get(row, "amt_s"))
        amt_a = _to_num(get(row, "amt_a"))
        diff = _to_num(get(row, "diff"))
        first_col = str(row[0]).strip() if row and row[0] is not None else ""

        # 首列是中文数字序号（一/二/三…）且无单位工程列时，视为分组标题
        is_group_row = False
        gname = None
        if CN_NUM.match(first_col) and not re.match(r"^\d+$", first_col):
            gname, is_group_row = name, True
            if amt_s is not None or amt_a is not None or diff is not None:
                warnings.append(f"分组行“{gname}”带有汇总金额，已按分组标题处理、未计入明细，"
                                "配平时请以其下明细合计为准。")

        if is_group_row:
            current = {"unit": gname or name, "items": []}
            groups.append(current)
            continue
        if not groups:
            groups.append(current)

        qty_s = _to_num(get(row, "qty_s"))
        qty_a = _to_num(get(row, "qty_a"))
        price_s = _to_num(get(row, "price_s"))
        price_a = _to_num(get(row, "price_a"))

        if amt_s is not None and amt_a is not None:
            net = amt_s - amt_a  # 净核减，正=核减
        elif diff is not None:
            net = diff
        else:
            continue  # 无金额信息，跳过

        current["items"].append({
            "name": name,
            "amt_s": amt_s, "amt_a": amt_a,
            "qty_s": qty_s, "qty_a": qty_a,
            "price_s": price_s, "price_a": price_a,
            "net": round(net, 2),
            "reason": infer_reason(qty_s, qty_a, price_s, price_a),
            "note": str(get(row, "note") or "").strip(),
            "excluded": False,
        })

    groups = [g for g in groups if g["items"]]
    if not groups:
        return None
    return {"sheet": ws.title, "columns": mapping, "groups": groups, "warnings": warnings}


def parse_workbook(path):
    """解析工作簿中所有合格的工作表，返回 {'sheets': [...]}。
    每个 sheet 结果含 sheet/columns/groups/warnings。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        try:
            res = _parse_jianming_sheet(ws)  # 优先识别"工程审核简明表"版式
            if res is None:
                res = _parse_sheet(ws)       # 通用关键词版式
        except Exception:
            res = None
        if res:
            sheets.append(res)
    wb.close()
    if not sheets:
        raise ValueError("未能识别对比表结构：请确认表中含有“项目名称、送审金额、审核金额”等栏目。")
    return {"sheets": sheets}


def infer_reason(qty_s, qty_a, price_s, price_a):
    """按 项目说明 规则推断调整原因。"""
    qty_changed = qty_s is not None and qty_a is not None and abs(qty_s - qty_a) > 1e-9
    price_changed = price_s is not None and price_a is not None and abs(price_s - price_a) > 1e-9
    if qty_changed and price_changed:
        return "工程量及单价有误"
    if qty_changed:
        return "工程量有误"
    if price_changed:
        return "单价有误"
    if qty_s is None and price_s is None:
        return "待确认"
    return "待确认"


REASON_OPTIONS = ["工程量有误", "单价有误", "工程量及单价有误", "取费标准有误", "待确认"]


def compose_text(sheets, threshold=1000.0, total_net=None, overrides=None, excluded_sheets=None):
    """组稿。sheets: parse_workbook 返回的 sheets 列表。
    overrides: {"表序号,组序号,条目序号": {'reason':..., 'excluded':...}}
    excluded_sheets: 不采用的工作表序号集合。
    total_net: 总净核减（正=核减）；为 None 时不做配平。
    返回 {'text', 'listed_sum', 'remainder', 'skipped'}
    """
    lines = []
    listed_sum = 0.0
    skipped = 0
    for s_idx, sheet in enumerate(sheets):
        if excluded_sheets and s_idx in excluded_sheets:
            continue
        for g_idx, g in enumerate(sheet["groups"]):
            group_lines = []
            for i, it in enumerate(g["items"]):
                ov = (overrides or {}).get(f"{s_idx},{g_idx},{i}", {})
                reason = ov.get("reason", it["reason"])
                excluded = ov.get("excluded", it.get("excluded", False))
                net = it["net"]
                if excluded or abs(net) < threshold:
                    skipped += 1
                    continue
                listed_sum += net
                sign = "负" if net > 0 else "正"  # 净核减为正→负偏差（核减）
                amt = _fmt(abs(net))
                sent = f"送审{_fmt(it['amt_s'])}元" if it["amt_s"] is not None else "经审核"
                if reason == "待确认":
                    line = f"{it['name']}：{sent}，审核依据不足，此项暂按{sign}偏差{amt}元计列，待进一步确认。"
                else:
                    line = f"{it['name']}：{sent}，由于{reason}，审核予以调整，此项{sign}偏差{amt}元。"
                group_lines.append(line)
            if group_lines:
                lines.append(f"{g['unit']}：")
                lines.extend(group_lines)

    remainder = None
    if total_net is not None:
        remainder = round(total_net - listed_sum, 2)
        if abs(remainder) >= 0.005:
            if remainder > 0:
                tail = f"由于取费基数变化，相应费率金额发生改变，此部分共净核减{_fmt(remainder)}元。"
            else:
                tail = f"由于取费基数变化，相应费率金额发生改变，此部分共净核增{_fmt(abs(remainder))}元。"
            lines.append(tail)

    return {
        "text": "\n".join(lines),
        "listed_sum": round(listed_sum, 2),
        "remainder": remainder,
        "skipped": skipped,
    }
