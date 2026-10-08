# -*- coding: utf-8 -*-
"""工程结算审核报告生成软件 - 本地服务端
双击启动后在本机 127.0.0.1 提供服务，浏览器即界面，不联网。
"""
import json
import os
import re
import subprocess
import sys
import threading
import webbrowser

from flask import Flask, jsonify, request, send_file, send_from_directory

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECTS_DIR = os.path.join(APP_DIR, "projects")
TEMPLATES_DIR = os.path.join(APP_DIR, "templates")
CONFIG_PATH = os.path.join(APP_DIR, "config.json")
DEFAULT_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates", "结算审核模板_占位版.docx")
DEFAULT_OUT_DIR = os.path.join(BASE, "outputs", "生成报告")
os.makedirs(PROJECTS_DIR, exist_ok=True)
os.makedirs(TEMPLATES_DIR, exist_ok=True)

sys.path.insert(0, APP_DIR)
import report

app = Flask(__name__, static_folder=APP_DIR)


def _safe_name(name):
    return re.sub(r'[\\/:*?"<>|]', "_", str(name).strip()) or "未命名项目"


def _load_config():
    cfg = {"out_dir": DEFAULT_OUT_DIR, "active_template": ""}
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, encoding="utf-8") as f:
                cfg.update(json.load(f))
        except Exception:
            pass
    return cfg


def _save_config(cfg):
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


def _active_template_path():
    cfg = _load_config()
    name = cfg.get("active_template", "")
    if name:
        p = os.path.join(TEMPLATES_DIR, name)
        if os.path.exists(p):
            return p
    return DEFAULT_TEMPLATE


# ---------------- 页面与字段 ----------------

@app.route("/")
def index():
    return send_from_directory(APP_DIR, "index.html")


@app.route("/api/fields")
def api_fields():
    with open(os.path.join(APP_DIR, "fields.json"), encoding="utf-8") as f:
        return jsonify(json.load(f))


# ---------------- 项目保存/载入 ----------------

@app.route("/api/projects", methods=["GET"])
def api_list_projects():
    items = [fn[:-5] for fn in sorted(os.listdir(PROJECTS_DIR)) if fn.endswith(".json")]
    return jsonify(items)


@app.route("/api/projects/<name>", methods=["GET"])
def api_load_project(name):
    path = os.path.join(PROJECTS_DIR, _safe_name(name) + ".json")
    if not os.path.exists(path):
        return jsonify({"error": "项目不存在"}), 404
    with open(path, encoding="utf-8") as f:
        return jsonify(json.load(f))


@app.route("/api/projects/<name>", methods=["POST"])
def api_save_project(name):
    data = request.get_json(force=True)
    path = os.path.join(PROJECTS_DIR, _safe_name(name) + ".json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return jsonify({"ok": True, "name": _safe_name(name)})


# ---------------- 保存路径 ----------------

@app.route("/api/settings", methods=["GET"])
def api_get_settings():
    cfg = _load_config()
    return jsonify({"out_dir": cfg["out_dir"], "active_template": cfg.get("active_template", "")})


@app.route("/api/settings/out_dir", methods=["POST"])
def api_set_out_dir():
    data = request.get_json(force=True)
    out_dir = str(data.get("out_dir", "")).strip()
    if not out_dir:
        return jsonify({"error": "路径不能为空"}), 400
    try:
        os.makedirs(out_dir, exist_ok=True)
    except Exception as e:
        return jsonify({"error": f"路径无效：{e}"}), 400
    cfg = _load_config()
    cfg["out_dir"] = out_dir
    _save_config(cfg)
    return jsonify({"ok": True, "out_dir": out_dir})


@app.route("/api/browse_folder", methods=["POST"])
def api_browse_folder():
    """弹出系统文件夹选择框（本机操作），返回所选路径。"""
    code = (
        "import tkinter as tk;from tkinter import filedialog;"
        "r=tk.Tk();r.withdraw();r.attributes('-topmost',True);"
        "p=filedialog.askdirectory(title='选择报告保存位置');"
        "r.destroy();print(p)"
    )
    try:
        res = subprocess.run([sys.executable, "-c", code],
                             capture_output=True, text=True, timeout=120)
        path = res.stdout.strip().replace("/", "\\")
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    return jsonify({"path": path})


# ---------------- 模板管理 ----------------

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}?")

def _known_field_names():
    with open(os.path.join(APP_DIR, "fields.json"), encoding="utf-8") as f:
        return {x["name"] for x in json.load(f)}


def _scan_template(path):
    """提取模板中的占位符，返回 (字段集合, 未闭合数量)。"""
    from docx import Document
    doc = Document(path)
    texts = [p.text for p in report._all_paragraphs(doc)]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    full = "\n".join(texts)
    found = set(re.findall(r"\{\{([^{}]+)\}\}", full))
    unclosed = full.count("{{") - len(re.findall(r"\{\{[^{}]*\}\}", full))
    return found, unclosed


@app.route("/api/templates", methods=["GET"])
def api_list_templates():
    cfg = _load_config()
    items = [fn for fn in sorted(os.listdir(TEMPLATES_DIR)) if fn.endswith(".docx")]
    return jsonify({
        "builtin": os.path.basename(DEFAULT_TEMPLATE),
        "imported": items,
        "active": cfg.get("active_template", "") or "(内置) " + os.path.basename(DEFAULT_TEMPLATE),
    })


@app.route("/api/templates/import", methods=["POST"])
def api_import_template():
    f = request.files.get("file")
    if not f or not f.filename.lower().endswith(".docx"):
        return jsonify({"error": "请上传 .docx 格式的模板文件"}), 400
    name = _safe_name(os.path.splitext(f.filename)[0])
    path = os.path.join(TEMPLATES_DIR, name + ".docx")
    f.save(path)

    warnings = []
    # 先尝试自动埋点：生成 占位版 并优先启用
    embedded_name = name + "_占位版.docx"
    embedded_path = os.path.join(TEMPLATES_DIR, embedded_name)
    embedded_count, missed = 0, []
    try:
        import embed
        embedded_path, embedded_count, missed = embed.auto_embed(path, embedded_path)
    except Exception as e:
        warnings.append(f"自动埋点失败：{e}，已保留原模板。")

    use_name = name + ".docx"
    if embedded_count >= 10:
        use_name = embedded_name
        if missed:
            warnings.append("以下位置未能自动识别，请在模板中手动改为 {{字段名}}：" + "、".join(missed))
    else:
        if os.path.exists(embedded_path):
            os.remove(embedded_path)
        warnings.append("未能从模板中自动识别出足够的可变内容。请把模板里的可变内容"
                        "（工程名称、单位、金额等）手动改为 {{字段名}} 后重新导入，"
                        "或在模板列表切换回内置模板。")

    # 校验最终启用的模板
    try:
        found, unclosed = _scan_template(os.path.join(TEMPLATES_DIR, use_name))
        unknown = sorted(found - _known_field_names())
        if unknown:
            warnings.append("模板中存在软件不认识的字段：" + "、".join("{{%s}}" % u for u in unknown) + "，生成时将原样保留。")
        if unclosed:
            warnings.append(f"检测到 {unclosed} 处未闭合的 '{{{{'，请检查模板。")
    except Exception as e:
        return jsonify({"error": f"模板读取失败：{e}"}), 400

    cfg = _load_config()
    cfg["active_template"] = use_name
    _save_config(cfg)
    return jsonify({"ok": True, "name": use_name, "warnings": warnings,
                    "fields_found": embedded_count or len(found),
                    "auto_embedded": use_name == embedded_name})


@app.route("/api/templates/active", methods=["POST"])
def api_set_active_template():
    data = request.get_json(force=True)
    name = str(data.get("name", "")).strip()
    if name and not os.path.exists(os.path.join(TEMPLATES_DIR, name)):
        return jsonify({"error": "模板不存在"}), 404
    cfg = _load_config()
    cfg["active_template"] = name
    _save_config(cfg)
    return jsonify({"ok": True})


# ---------------- Excel 对比表 ----------------

UPLOADS_DIR = os.path.join(APP_DIR, "uploads")
os.makedirs(UPLOADS_DIR, exist_ok=True)
_EXCEL_CACHE = {}  # token -> parsed result


@app.route("/api/excel/upload", methods=["POST"])
def api_excel_upload():
    f = request.files.get("file")
    if not f or not f.filename.lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"error": "请上传 .xlsx 格式的对比表"}), 400
    import excel_import
    import time
    token = _safe_name(os.path.splitext(f.filename)[0]) + "_" + str(int(time.time()))
    path = os.path.join(UPLOADS_DIR, token + ".xlsx")
    f.save(path)
    try:
        parsed = excel_import.parse_workbook(path)
    except Exception as e:
        os.remove(path)
        return jsonify({"error": str(e)}), 400
    _EXCEL_CACHE.clear()
    _EXCEL_CACHE[token] = parsed
    summary = []
    for s_idx, sh in enumerate(parsed["sheets"]):
        summary.append({
            "sheet": sh["sheet"],
            "warnings": sh["warnings"],
            "groups": [{"unit": g["unit"], "count": len(g["items"])} for g in sh["groups"]],
        })
    return jsonify({"ok": True, "token": token, "sheets": parsed["sheets"], "summary": summary})


@app.route("/api/excel/compose", methods=["POST"])
def api_excel_compose():
    data = request.get_json(force=True)
    token = data.get("token", "")
    parsed = _EXCEL_CACHE.get(token)
    if not parsed:
        return jsonify({"error": "对比表数据已失效，请重新上传"}), 400
    import excel_import
    threshold = float(data.get("threshold") or 1000)
    total_raw = str(data.get("total_net", "")).replace(",", "").strip()
    total_net = float(total_raw) if total_raw else None
    result = excel_import.compose_text(
        parsed["sheets"],
        threshold=threshold,
        total_net=total_net,
        overrides=data.get("overrides") or {},
        excluded_sheets=set(data.get("excluded_sheets") or []),
    )
    return jsonify({"ok": True, **result})


# ---------------- 生成与下载 ----------------

@app.route("/api/generate", methods=["POST"])
def api_generate():
    values = request.get_json(force=True)
    out_dir = str(values.pop("_out_dir", "")).strip() or _load_config()["out_dir"]
    tpl = _active_template_path()
    # 防护：模板不含任何占位符时生成结果必然与模板一致，直接拦截并说明
    try:
        found, _ = _scan_template(tpl)
    except Exception as e:
        return jsonify({"error": f"模板读取失败：{e}"}), 500
    if not found:
        return jsonify({"error": "当前模板中没有任何 {{字段}} 占位符，生成结果只会是模板原文。"
                                 "请在“报告模板”卡片切换回内置模板或带占位符的模板；"
                                 "若需使用此模板，请先把可变内容改为 {{字段名}} 再导入。"}), 400
    try:
        os.makedirs(out_dir, exist_ok=True)
        path, warnings = report.generate(values, template_path=tpl, out_dir=out_dir)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    return jsonify({"ok": True, "file": os.path.basename(path),
                    "dir": os.path.dirname(path), "warnings": warnings})


@app.route("/api/download", methods=["POST"])
def api_download():
    data = request.get_json(force=True)
    path = os.path.join(str(data.get("dir", "")), str(data.get("file", "")))
    if not os.path.exists(path):
        return jsonify({"error": "文件不存在"}), 404
    return send_file(path, as_attachment=True)


@app.route("/api/open", methods=["POST"])
def api_open():
    """用系统默认程序打开生成的报告（本机操作）。"""
    data = request.get_json(force=True)
    path = os.path.join(str(data.get("dir", "")), str(data.get("file", "")))
    if not os.path.exists(path):
        return jsonify({"error": "文件不存在"}), 404
    os.startfile(path)
    return jsonify({"ok": True})


def main():
    port = 8765
    url = f"http://127.0.0.1:{port}"
    threading.Timer(1.0, lambda: webbrowser.open(url)).start()
    print(f"工程结算审核报告生成软件已启动：{url}")
    print("关闭本窗口即退出软件。")
    app.run(host="127.0.0.1", port=port, debug=False)


if __name__ == "__main__":
    main()
