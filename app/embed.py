# -*- coding: utf-8 -*-
"""模板自动埋点引擎：把一份结算审核报告 docx 自动转为带 {{字段}} 占位符的模板。
规则基于报告体例（标签/正则定位），与具体项目内容无关，适用于同类报告模板。
所有替换跨 run 进行、保留格式；含图形的 run（横线等，包括 mc:AlternateContent 包裹）一律不动。
返回 (输出路径, 已埋字段数, 未命中规则列表)。
"""
import copy
import os
import re

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

PLACEHOLDER_RE = re.compile(r"\{\{[^{}]+\}\}")


# ---------------- 基础工具 ----------------

MC_URI = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def run_has_drawing(r):
    el = r._element
    for tag in ("w:pict", "w:drawing", "w:object"):
        if el.find(".//" + qn(tag)) is not None:
            return True
    return el.find(".//{%s}AlternateContent" % MC_URI) is not None


def _set_run_text(r, text):
    if run_has_drawing(r):
        return
    r.text = text


def replace_in_paragraph(p, old, new):
    """跨 run 替换 old->new；新文字继承匹配起始 run 格式；图形 run 不动。"""
    runs = p.runs
    count = 0
    while True:
        full = "".join(r.text for r in runs)
        idx = full.find(old)
        if idx == -1:
            break
        end = idx + len(old)
        pos = 0
        for r in runs:
            rlen = len(r.text)
            rs, re_ = pos, pos + rlen
            pos = re_
            if re_ <= idx or rs >= end:
                continue
            if run_has_drawing(r):
                continue
            local_s = max(idx - rs, 0)
            local_e = min(end - rs, rlen)
            if rs <= idx:
                r.text = r.text[:local_s] + new + (r.text[local_e:] if re_ >= end else "")
            else:
                r.text = r.text[local_e:] if re_ >= end else ""
        count += 1
    return count


def set_paragraph_text(p, new):
    """整段设为模板文本（保留第一个非图形 run 的格式；图形 run 不动）。
    注意：会丢弃段内原有前导空格，仅用于整段替换场景。"""
    target = None
    for r in p.runs:
        if not run_has_drawing(r):
            target = r
            break
    if target is None:
        p.add_run(new)
        return
    target.text = new
    for r in p.runs:
        if r._element is not target._element:
            _set_run_text(r, "")


def _delete(el):
    el.getparent().remove(el)


# ---------------- 埋点规则 ----------------

AMOUNTS_TPL = ("本工程建设单位审核价为{{送审金额}}元，审核金额为{{审定金额}}元，"
               "核增{{核增金额}}元，核减{{核减金额}}元，净核减{{净核减金额}}元。修正情况分析如下：")
CONCLUSION_TPL = ("本工程建设单位审定金额为{{送审金额}}元；审核金额为{{审定金额}}元；"
                  "核增金额为{{核增金额}}元；核减金额为{{核减金额}}元；净核减(增)金额为{{净核减金额}}元。")


def auto_embed(src_path, dst_path):
    doc = Document(src_path)
    paras = doc.paragraphs
    missed = []
    embedded = set()

    def find(pred, reverse=False):
        seq = reversed(paras) if reverse else paras
        for p in seq:
            try:
                if pred(p.text):
                    return p
            except Exception:
                continue
        return None

    def mark(name, ok):
        if ok:
            embedded.add(name)
        else:
            missed.append(name)

    # 1) 文头咨询公司：第一个非空段
    first = next((p for p in paras if p.text.strip()), None)
    company = first.text.strip() if first else ""
    if first is not None:
        set_paragraph_text(first, "{{咨询公司}}")
        embedded.add("咨询公司")
    else:
        missed.append("咨询公司")

    # 2) 文号行：含 〔xxxx〕xx号- 的段落
    p = find(lambda t: re.search(r"〔\s*\d{4}\s*〕\s*\d+\s*号\s*-", t))
    mark("报告文号", p is not None)
    if p is not None:
        set_paragraph_text(p, "{{报告文号}}-{{工程名称}}")

    # 3) 关于X的：紧邻"结算审核报告"标题之前的段落
    idx_title = next((i for i, p in enumerate(paras) if p.text.strip() == "结算审核报告"), None)
    if idx_title and idx_title > 0:
        set_paragraph_text(paras[idx_title - 1], "关于{{工程名称}}的")
        embedded.add("工程名称")
    else:
        missed.append("关于…的标题行")

    # 4) 主送单位：首段（我公司接受…）之前、以"："结尾的段落
    idx_open = next((i for i, p in enumerate(paras) if p.text.strip().startswith("我公司接受")), None)
    if idx_open:
        prev = paras[idx_open - 1] if idx_open > 0 else None
        if prev is not None and prev.text.strip().endswith("："):
            set_paragraph_text(prev, "{{委托单位}}：")
            embedded.add("委托单位")
        # 5) 首段正文：对由X负责建设的Y结算进行了审核。X的责任是…
        p = paras[idx_open]
        m = re.search(r"对由(.+?)负责建设的(.+?)结算进行了审核[。．]?(.+?)的责任是", p.text)
        if m:
            replace_in_paragraph(p, m.group(1), "{{建设单位}}")
            replace_in_paragraph(p, m.group(2), "{{工程名称}}")
            # 责任主体（group3 与 group1 通常相同，已被第一步替换；若不同再替换）
            if m.group(3) != m.group(1):
                replace_in_paragraph(p, m.group(3), "{{建设单位}}")
            embedded.update(["建设单位", "工程名称"])
        else:
            missed.append("首段建设单位/工程名称")

    # 6) 项目概况（标签定位）
    label_rules = [
        ("批复总投资：", "批复总投资：{{批复总投资}}；", "批复总投资"),
        ("建设单位：", "建设单位：{{建设单位}}；", "建设单位"),
        ("代建单位：", "代建单位：{{代建单位}}；", "代建单位"),
        ("施工单位：", "施工单位：{{施工单位}}；", "施工单位"),
        ("监理单位：", "监理单位：{{监理单位}}；", "监理单位"),
    ]
    for anchor, tpl, name in label_rules:
        p = find(lambda t, a=anchor: t.strip().startswith(a))
        mark(name, p is not None)
        if p is not None:
            set_paragraph_text(p, tpl)

    p = find(lambda t: t.strip().startswith("本工程合同额"))
    mark("合同额/工程内容", p is not None)
    if p is not None:
        set_paragraph_text(p, "本工程合同额{{合同额}}元。工程内容：{{工程内容}}；")

    p = find(lambda t: t.strip().startswith("本工程合同总工期"))
    mark("工期说明", p is not None)
    if p is not None:
        set_paragraph_text(p, "{{工期说明}}")

    # 7) 审核依据：连续以《开头的段落合并为一个占位段（止于非《开头段）
    first_book = next((i for i, p in enumerate(paras) if p.text.strip().startswith("《")), None)
    if first_book is not None:
        j = first_book
        while j < len(paras) and paras[j].text.strip().startswith("《"):
            j += 1
        set_paragraph_text(paras[first_book], "{{审核依据}}")
        body = doc.element.body
        for k in range(first_book + 1, j):
            _delete(paras[k]._element)
        embedded.add("审核依据")
    else:
        missed.append("审核依据")

    # 8) 审核内容/程序
    p = find(lambda t: t.strip().startswith("工程量、联系单"))
    mark("审核内容", p is not None)
    if p is not None:
        set_paragraph_text(p, "{{审核内容}}")
    p = find(lambda t: t.strip().startswith("审核程序："))
    mark("审核程序", p is not None)
    if p is not None:
        set_paragraph_text(p, "{{审核程序}}")

    # 9) 有关事项说明
    p = find(lambda t: "为基期价格" in t)
    if p is not None:
        m = re.search(r"本工程以(.+?)为基期价格", p.text)
        if m:
            replace_in_paragraph(p, m.group(1), "{{基期价格}}")
            embedded.add("基期价格")
        else:
            missed.append("基期价格")
    p = find(lambda t: t.strip().startswith("工程下浮率"))
    mark("下浮率", p is not None)
    if p is not None:
        set_paragraph_text(p, "工程下浮率{{下浮率}}；")
    p = find(lambda t: t.strip().startswith("工期延期"))
    if p is not None:
        set_paragraph_text(p, "{{工期延期说明}}")
        embedded.add("工期延期说明")
    p = find(lambda t: t.strip().startswith("施工单位人员到位率"))
    if p is not None:
        set_paragraph_text(p, "{{人员到位说明}}")
        embedded.add("人员到位说明")
    p = find(lambda t: "创标化工地" in t)
    if p is not None:
        set_paragraph_text(p, "{{标化工地说明}}")
        embedded.add("标化工地说明")

    # 10) 金额段 + 核减明细合并
    p = find(lambda t: "审核价为" in t and "净核减" in t)
    mark("金额五字段", p is not None)
    if p is not None:
        set_paragraph_text(p, AMOUNTS_TPL)
        p_end = find(lambda t: t.strip() == "审核结论")
        if p_end is not None:
            body = doc.element.body
            children = list(body)
            si = children.index(p._element)
            ei = children.index(p_end._element)
            mid = children[si + 1:ei]
            if mid:
                set_paragraph_text(Paragraph(mid[0], p._parent), "{{核减明细}}")
                for el in mid[1:]:
                    _delete(el)
                embedded.add("核减明细")
        else:
            missed.append("核减明细（未找到“审核结论”标题）")

    p = find(lambda t: "审定金额为" in t and "净核减(增)" in t)
    if p is not None:
        set_paragraph_text(p, CONCLUSION_TPL)

    # 11) 落款：最后一个与公司名相同的段——只替换公司名，保留前导空格（右对齐效果）
    if company:
        p = find(lambda t: t.strip() == company, reverse=True)
        if p is not None and p is not first:
            replace_in_paragraph(p, company, "{{咨询公司}}")

    # 12) 落款日期：最后一个纯日期段——只替换日期文字，保留前导空格
    p = find(lambda t: re.match(r"^\s*\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日\s*$", t), reverse=True)
    if p is not None:
        m = re.search(r"\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日", p.text)
        if m:
            replace_in_paragraph(p, m.group(0), "{{报告日期}}")
            embedded.add("报告日期")

    # 13) 抄送：替换"抄送："之后的内容，保留拖尾下划线空格
    p = find(lambda t: t.strip().startswith("抄送："))
    if p is not None:
        m = re.match(r"^(\s*抄送：)(.+?)(\s*)$", p.text)
        if m and m.group(2).strip():
            replace_in_paragraph(p, m.group(2), "{{抄送单位}}")
            embedded.add("抄送单位")

    doc.save(dst_path)
    return dst_path, len(embedded), missed
